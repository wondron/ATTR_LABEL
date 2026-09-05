from __future__ import annotations

import asyncio
import io
import json
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

PACKAGE_ROOT = Path(__file__).resolve().parents[1]
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from annotation_app.config import MISSING_REVISION
from annotation_app.export import EXCEL_HEADERS, create_export_archive
from annotation_app.main import ExportFileResponse, create_app
from annotation_app.models import AnnotationValues
from annotation_app.repository import AnnotationRepository, ImageNotFoundError, RepositoryError


def make_image(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (8, 6), (210, 130, 40)).save(path)


def save_example(repository: AnnotationRepository, image_id: str, name: str = "包子") -> None:
    repository.save_annotation(
        image_id,
        AnnotationValues(
            food_name=name,
            food_count=2,
            quality=123.5,
            device_model="C9277A",
            container_type=["陶瓷容器", "油纸"],
            accessory_type=["烤盘"],
            rack_level=["1", "2"],
            food_size=8,
        ),
    )


class ExportDeleteTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / "当前 文件夹"
        self.root.mkdir()

    def client(self) -> TestClient:
        return TestClient(create_app(self.root, allowed_data_roots=(self.root.parent,)))

    def test_zip_contains_only_labeled_originals_and_nine_typed_excel_columns(self) -> None:
        for name in ("子目录/同名.jpg", "另一个目录/同名.jpg", "未标注.jpg", "=1+1.jpg"):
            make_image(self.root / name)
        repository = AnnotationRepository(self.root)
        for name in ("子目录/同名.jpg", "另一个目录/同名.jpg", "=1+1.jpg"):
            save_example(repository, name, '=HYPERLINK("https://example.invalid","文字")')

        generated = []
        def capture(*args, **kwargs):
            archive = create_export_archive(*args, **kwargs)
            generated.append(archive)
            return archive

        with self.client() as client, patch("annotation_app.main.create_export_archive", side_effect=capture):
            response = client.get("/api/v1/export", params={"directory_generation": 0})
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertIn("当前 文件夹.zip", unquote(response.headers["content-disposition"]))
        self.assertFalse(generated[0].path.parent.exists(), "下载完成必须清理临时文件")

        with ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(set(archive.namelist()), {
                "子目录/同名.jpg", "另一个目录/同名.jpg", "=1+1.jpg", "标注数据.xlsx",
            })
            for name in archive.namelist():
                self.assertFalse(name.startswith("/"))
                self.assertNotIn("..", Path(name).parts)
                if name.endswith(".jpg"):
                    self.assertEqual(archive.read(name), (self.root / name).read_bytes())
            xlsx = archive.read("标注数据.xlsx")
        workbook = load_workbook(io.BytesIO(xlsx), data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), EXCEL_HEADERS)
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.max_column, 9)
        for row in sheet.iter_rows(min_row=2):
            self.assertEqual(row[1].data_type, "s")
            self.assertTrue(row[1].value.startswith("=HYPERLINK"))
            self.assertEqual([cell.value for cell in row[2:]], [
                2, 123.5, "C9277A", "陶瓷容器、油纸", "烤盘", "1、2", 8,
            ])
        with ZipFile(io.BytesIO(xlsx)) as contents:
            self.assertNotIn(b"<f>", contents.read("xl/worksheets/sheet1.xml"))

    def test_no_saved_annotations_returns_exact_empty_message(self) -> None:
        make_image(self.root / "未标注.jpg")
        with self.client() as client:
            response = client.get("/api/v1/export", params={"directory_generation": 0})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()["detail"]["message"], "暂无已标注数据可下载")

    def test_corrupt_annotation_stops_export_and_cleans_temporary_directory(self) -> None:
        make_image(self.root / "损坏.jpg")
        (self.root / "损坏.json").write_text("{broken", encoding="utf-8")
        created = []
        real_temporary = tempfile.TemporaryDirectory
        def temporary(*args, **kwargs):
            directory = real_temporary(*args, **kwargs)
            created.append(Path(directory.name))
            return directory
        with self.client() as client, patch("annotation_app.export.TemporaryDirectory", side_effect=temporary):
            response = client.get("/api/v1/export", params={"directory_generation": 0})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["detail"]["code"], "invalid_annotation_file")
        self.assertIn("损坏.json", response.json()["detail"]["message"])
        self.assertTrue(created)
        self.assertTrue(all(not path.exists() for path in created))

    def test_original_changes_during_export_returns_error(self) -> None:
        image = self.root / "变化.jpg"
        make_image(image)
        repository = AnnotationRepository(self.root)
        save_example(repository, image.name)
        original_write = ZipFile.write
        def changed_write(archive, filename, *args, **kwargs):
            result = original_write(archive, filename, *args, **kwargs)
            if Path(filename) == image:
                with image.open("ab") as stream:
                    stream.write(b"changed")
            return result
        with patch.object(ZipFile, "write", changed_write):
            with self.assertRaisesRegex(RepositoryError, "图像仍在写入或已经变化"):
                create_export_archive(repository)

    def test_response_send_failure_cleans_archive(self) -> None:
        make_image(self.root / "a.jpg")
        repository = AnnotationRepository(self.root)
        save_example(repository, "a.jpg")
        archive = create_export_archive(repository)
        response = ExportFileResponse(archive)
        async def receive():
            return {"type": "http.disconnect"}
        async def send(message):
            raise ConnectionError("client left")
        with self.assertRaises(ConnectionError):
            asyncio.run(response({"type": "http", "method": "GET", "headers": []}, receive, send))
        self.assertFalse(archive.path.parent.exists())

    def test_delete_labeled_and_unlabeled_images_updates_listing_and_statistics(self) -> None:
        make_image(self.root / "a.jpg")
        make_image(self.root / "b.jpg")
        save_example(AnnotationRepository(self.root), "a.jpg")
        with self.client() as client:
            for name, total in (("a.jpg", 1), ("b.jpg", 0)):
                response = client.delete("/api/v1/image", params={"image_id": name, "directory_generation": 0})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json(), {
                    "status": "deleted", "image_id": name, "image_exists": False,
                    "annotation_exists": False, "directory_generation": 0,
                })
                listing = client.get("/api/v1/images", params={"directory_generation": 0}).json()
                statistics = client.get("/api/v1/statistics", params={"directory_generation": 0}).json()
                self.assertEqual(listing["summary"]["total"], total)
                self.assertEqual(statistics["summary"]["total"], total)
                self.assertEqual(statistics["summary"]["labeled"], 0)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_stale_generation_and_traversal_cannot_delete_or_export(self) -> None:
        make_image(self.root / "a.jpg")
        save_example(AnnotationRepository(self.root), "a.jpg")
        with self.client() as client:
            stale = client.delete("/api/v1/image", params={"image_id": "a.jpg", "directory_generation": 1})
            stale_export = client.get("/api/v1/export", params={"directory_generation": 1})
            self.assertEqual(stale.status_code, 409)
            self.assertEqual(stale_export.status_code, 409)
            for image_id in ("../a.jpg", "C:/a.jpg", "/a.jpg", "a.jpg:stream"):
                response = client.delete("/api/v1/image", params={"image_id": image_id, "directory_generation": 0})
                self.assertEqual(response.status_code, 400, response.text)
        self.assertTrue((self.root / "a.jpg").is_file())
        self.assertTrue((self.root / "a.json").is_file())

    def test_shared_sidecar_rejects_deletion_without_changing_either_image(self) -> None:
        make_image(self.root / "a.jpg")
        make_image(self.root / "a.png")
        save_example(AnnotationRepository(self.root), "a.jpg")
        with self.client() as client:
            response = client.delete("/api/v1/image", params={"image_id": "a.jpg", "directory_generation": 0})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["detail"]["code"], "shared_annotation")
        self.assertTrue(response.json()["detail"]["image_exists"])
        self.assertTrue(response.json()["detail"]["annotation_exists"])
        self.assertEqual(len(list(self.root.iterdir())), 3)

    def test_delete_failures_report_disk_state_after_each_possible_failed_step(self) -> None:
        original_unlink = Path.unlink
        for failed_suffix in (".jpg", ".json"):
            with self.subTest(failed_suffix=failed_suffix):
                image = self.root / "a.jpg"
                make_image(image)
                repository = AnnotationRepository(self.root)
                if not image.with_suffix(".json").exists():
                    save_example(repository, image.name)
                def fail_unlink(path, *args, **kwargs):
                    if path == image.with_suffix(failed_suffix):
                        raise PermissionError("simulated permission failure")
                    return original_unlink(path, *args, **kwargs)
                with self.client() as client, patch.object(Path, "unlink", fail_unlink):
                    response = client.delete("/api/v1/image", params={"image_id": "a.jpg", "directory_generation": 0})
                    self.assertEqual(response.status_code, 500, response.text)
                    detail = response.json()["detail"]
                    self.assertIn("simulated permission failure", detail["message"])
                    self.assertEqual(detail["image_exists"], image.exists())
                    self.assertEqual(detail["annotation_exists"], image.with_suffix(".json").exists())
                    self.assertTrue(detail["annotation_exists"])
                    self.assertEqual(detail["image_exists"], failed_suffix == ".jpg")
                    listing = client.get("/api/v1/images", params={"directory_generation": 0}).json()
                    self.assertEqual(listing["summary"]["total"], int(image.exists()))

    def test_symbolic_image_and_sidecar_links_cannot_delete_targets(self) -> None:
        target = self.root.parent / "outside.jpg"
        make_image(target)
        link = self.root / "link.jpg"
        try:
            link.symlink_to(target)
        except OSError:
            self.skipTest("当前 Windows 用户没有创建符号链接的权限")
        with self.client() as client:
            response = client.delete("/api/v1/image", params={"image_id": link.name, "directory_generation": 0})
            self.assertEqual(response.status_code, 400)
            self.assertTrue(target.is_file())
            image = self.root / "a.jpg"
            make_image(image)
            external_json = self.root.parent / "outside.json"
            external_json.write_text("{}", encoding="utf-8")
            image.with_suffix(".json").symlink_to(external_json)
            response = client.delete("/api/v1/image", params={"image_id": image.name, "directory_generation": 0})
            self.assertEqual(response.status_code, 500)
            self.assertTrue(image.is_file())
            self.assertTrue(external_json.is_file())

    def test_save_waits_for_delete_and_cannot_recreate_annotation_after_image_is_gone(self) -> None:
        make_image(self.root / "a.jpg")
        repository = AnnotationRepository(self.root)
        save_example(repository, "a.jpg")
        _, revision = repository.read_annotation("a.jpg")
        deletion_started = threading.Event()
        allow_delete = threading.Event()
        original_unlink = Path.unlink
        def paused_unlink(path, *args, **kwargs):
            if path == self.root / "a.jpg":
                deletion_started.set()
                if not allow_delete.wait(timeout=5):
                    raise TimeoutError("test release did not arrive")
            return original_unlink(path, *args, **kwargs)
        with ThreadPoolExecutor(max_workers=2) as executor, patch.object(Path, "unlink", paused_unlink):
            deletion = executor.submit(repository.delete_image, "a.jpg")
            self.assertTrue(deletion_started.wait(timeout=5))
            saving = executor.submit(repository.save_annotation, "a.jpg", AnnotationValues(), revision)
            try:
                self.assertFalse(saving.done())
            finally:
                allow_delete.set()
            self.assertFalse(deletion.result(timeout=5)["image_exists"])
            with self.assertRaises(ImageNotFoundError):
                saving.result(timeout=5)
        self.assertFalse((self.root / "a.json").exists())

    def test_annotation_cache_reuses_unchanged_json_and_invalidates_external_change(self) -> None:
        make_image(self.root / "a.jpg")
        repository = AnnotationRepository(self.root)
        save_example(repository, "a.jpg")
        with patch.object(repository, "_read_uncached_document", wraps=repository._read_uncached_document) as reader:
            repository.list_images()
            repository.collect_annotation_documents()
            self.assertEqual(reader.call_count, 1)
            sidecar = self.root / "a.json"
            payload = json.loads(sidecar.read_text(encoding="utf-8"))
            payload["annotations"]["food_name"] = "更新后的名称"
            sidecar.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            document, _ = repository.read_annotation("a.jpg")
            self.assertEqual(document.annotations.food_name, "更新后的名称")
            self.assertEqual(reader.call_count, 2)
            sidecar.write_text("{broken", encoding="utf-8")
            self.assertFalse(repository.list_images()[0]["annotation_valid"])
            repository.list_images()
            self.assertEqual(reader.call_count, 3)


if __name__ == "__main__":
    unittest.main()
